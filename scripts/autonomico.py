"""Nivel autonómico: créditos u obligaciones de altos cargos (artículo 10) y personal eventual (artículo 11)
de cada comunidad, por unidad, más las retribuciones de los diputados autonómicos.

Lee datos/fuentes/<año>/<comunidad>/ y escribe datos/<año>/autonomico_unidades.csv.
Uso: py scripts/autonomico.py [año]
"""
import collections
import csv
import json
import re
import sys
import urllib.parse

import openpyxl
import pdfplumber

from comun import ANYO, FUENTES, SALIDAS, FaltaFuente, leer_json_url, metadatos, num, ruta

SALIDA = SALIDAS / "autonomico_unidades.csv"


# Cada extractor devuelve un Counter {(unidad, articulo[, tipo]): importe en euros}

def andalucia():
    ws = openpyxl.load_workbook(ruta("andalucia", "gastos.xlsx"), read_only=True).worksheets[0]
    c = collections.Counter()
    filas = ws.iter_rows(values_only=True)
    cab = next(filas)
    i_sec, i_art, i_imp = cab.index("Denominación Sección"), cab.index("Artículo"), cab.index("Crédito Inicial")
    for r in filas:
        if r[i_art] in ("10", "11"):
            c[(r[i_sec].strip(), r[i_art])] += float(r[i_imp] or 0)
    return c


def aragon():
    nombres = {}
    with open(ruta("aragon", "estructura_organica.csv"), encoding="latin-1") as f:
        for r in csv.DictReader(f, delimiter=";"):
            nombres.setdefault(r["CENTRO GESTOR"], r["DESCRIPCION LARGA"].strip())
    c = collections.Counter()
    with open(ruta("aragon", "ejecucion_gastos.csv"), encoding="latin-1") as f:
        for r in csv.DictReader(f, delimiter=";"):
            art = r["Económico"][:2]
            if art in ("10", "11"):
                cg = r["Centro Gestor"]
                # La sección 01 agrupa Cortes, Justicia y Cámara de Cuentas: se separan por centro gestor
                unidad = nombres.get(cg, cg) if cg.startswith("01") else nombres.get(cg[:2] + "010", "Sección " + cg[:2])
                c[(unidad, art)] += num(r["Obligaciones Reconocidas"])
    return c


def canarias():
    c = collections.Counter()
    with open(ruta("canarias", "presupuesto_consolidado.csv"), encoding="latin-1") as f:
        for r in csv.DictReader(f, delimiter=";"):
            art = r["Subconcepto"][:2]
            if r["Capítulo"] == "1" and art in ("10", "11"):
                c[(r["Descripción Sección"].strip(), art)] += float(r["Importe"].replace(",", "."))
    return c


def castilla_leon():
    base = ("https://analisis.datosabiertos.jcyl.es/api/explore/v2.1/catalog/datasets/"
            "ejecucion-del-presupuesto-de-la-administracion-de-la-comunidad-gastos/records?")
    q = urllib.parse.urlencode({
        "select": "codigo_seccion,seccion,servicio,codigo_articulo,sum(obligado_4) as imp",
        "where": f"year(ano)={ANYO} and codigo_articulo in ('10','11')",
        "group_by": "codigo_seccion,seccion,servicio,codigo_articulo",
        "limit": 100,
    })
    d = leer_json_url(base + q, "castillayleon/api_art10_11.json")
    if not d.get("results"):
        (FUENTES / "castillayleon" / "api_art10_11.json").unlink()   # no guardar una respuesta vacía
        raise FaltaFuente(f"Castilla y León aún no publica la ejecución de {ANYO}")
    c = collections.Counter()
    for r in d["results"]:
        # La sección 20 agrupa Cortes e instituciones propias: se separan por servicio
        unidad = r["servicio"] if r["codigo_seccion"] == "20" else r["seccion"]
        c[(unidad.strip(), r["codigo_articulo"])] += r["imp"] or 0
    return c


def cataluna():
    base = "https://analisi.transparenciacatalunya.cat/resource/ajns-4mi7.json?"
    q = urllib.parse.urlencode({
        "$select": "secci_codi,secci,article_codi,sum(obligacions_reconegudes) as imp",
        "$where": f"exercici='{ANYO}' and mes='12 Desembre' and article_codi in ('10','11')",
        "$group": "secci_codi,secci,article_codi",
        "$limit": 500,
    })
    d = leer_json_url(base + q, "cataluna/api_art10_11.json")
    if not d:
        (FUENTES / "cataluna" / "api_art10_11.json").unlink()   # no guardar una respuesta vacía
        raise FaltaFuente(f"Cataluña aún no publica el cierre (diciembre) de {ANYO}")
    c = collections.Counter()
    for r in d:
        # Las secciones con código numérico son organismos autónomos y entes
        tipo = "organismo" if r["secci_codi"].isdigit() and r["secci_codi"] not in ("70", "75", "77") else None
        c[(r["secci"].strip(), r["article_codi"], tipo)] += float(r.get("imp", 0))
    return c


def navarra():
    nombres = {}
    with open(ruta("navarra", "presupuesto_gastos.csv"), encoding="utf-8", errors="replace") as f:
        for r in csv.reader(f):
            m = re.match(r"Total (\w\w)-(.+)", r[0])
            if m:
                nombres.setdefault(m.group(1), m.group(2).strip())
    c = collections.Counter()
    with open(ruta("navarra", "ejecucion_gastos.csv"), encoding="utf-8", errors="replace") as f:
        lector = csv.reader(f)
        next(lector)
        for r in lector:
            p = r[0].split()
            if len(p) < 4 or not re.fullmatch(r"\d{4}", p[2]):
                continue
            org, art, prog = p[0], p[2][:2], p[3]
            if art not in ("10", "11"):
                continue
            # En el Servicio Navarro de Salud el artículo 11 es personal sanitario temporal, no político
            if art == "11" and org[:2] in ("52", "53", "54") and prog.startswith("3"):
                continue
            dep = org[:1] + "0"
            c[(nombres.get(dep, "Departamento " + dep), art, "gobierno")] += num(r[3])
    # El Parlamento solo figura como transferencia: se toma de la auditoría de la Cámara de Comptos
    return c + desde_json("navarra", "parlamento.json")


def valencia():
    c = collections.Counter()
    fichero = ruta("valencia", "ejecucion_gastos.csv")
    # El fichero es acumulado mensual: se toma el último mes publicado
    with open(fichero, encoding="utf-8", errors="replace") as f:
        ultimo = max(int(r["ac_mes"]) for r in csv.DictReader(f, delimiter=";"))
    if ultimo < 12:
        raise FaltaFuente(f"Valencia solo publica la ejecución de {ANYO} hasta el mes {ultimo}")
    with open(fichero, encoding="utf-8", errors="replace") as f:
        for r in csv.DictReader(f, delimiter=";"):
            if int(r["ac_mes"]) != ultimo or r["cd_art"] not in ("G10", "G11"):
                continue
            unidad = r["ds_org_sec_c"].split("-", 1)[-1].strip()
            tipo = None if r["ds_org_ent_c"].strip().lower().startswith("generalitat") else "organismo"
            c[(unidad, r["cd_art"][1:], tipo)] += float((r["im_recon"] or "0").replace(",", "."))
    return c


def pdf_por_secciones(fichero, patron_seccion, factor=1):
    """Suma las líneas '10 ALTOS CARGOS <importe>' y '11 PERSONAL EVENTUAL... <importe>' de un estado
    de gastos en PDF, atribuyéndolas a la última cabecera de sección vista en la página."""
    art = re.compile(r"^(10|11) (ALTOS CARGOS|PERSONAL EVENTUAL)[A-ZÁÉÍÓÚ ]* ([\d.]+(?:,\d+)?)$")
    sec = re.compile(patron_seccion)
    c = collections.Counter()
    with pdfplumber.open(fichero) as pdf:
        actual = None
        for page in pdf.pages:
            lineas = (page.extract_text() or "").splitlines()
            for l in lineas[:6]:
                m = sec.match(l)
                if m:
                    actual = m.group(2).strip()
            for l in lineas:
                m = art.match(l.strip())
                if m and actual:
                    c[(actual, m.group(1))] += num(m.group(3)) * factor
    if not c:
        raise ValueError(f"no se reconoce el formato de {fichero.name}: revisar el PDF")
    return c


def murcia():
    return pdf_por_secciones(ruta("murcia", "presupuestos_I.pdf"), r"^Secci[oó]n: (\d+) (.+)$")


def asturias():
    return pdf_por_secciones(ruta("asturias", "tomo2_estado_gastos.pdf"), r"^Secci[oó]n (\d+) (.+)$")


def cantabria():
    return pdf_por_secciones(ruta("cantabria", "estado_ingresos_gastos.pdf"), r"^SECCI[OÓ]N (\d+) (.+)$")


def castilla_la_mancha():
    # Importes en miles de euros
    return pdf_por_secciones(ruta("castillalamancha", "tomo1.pdf"), r"^Secci[oó]n (\d+) (.+)$", factor=1000)


def madrid():
    """Libro 03, 'Resumen general por tipo de centro, artículos y secciones': tabla con una columna
    por sección. Las cifras se asignan a la columna cuyo encabezado queda más cerca por la derecha."""
    nombres = {}
    c = collections.Counter()
    control = collections.Counter()
    with pdfplumber.open(ruta("madrid", "libro03.pdf")) as pdf:
        for page in pdf.pages:
            texto = page.extract_text() or ""
            for l in texto.splitlines():
                m = re.match(r"^(\d\d)-([A-ZÁÉÍÓÚÑ ,.Y]+?)\s+[\d.]+", l)
                if m:
                    nombres.setdefault(m.group(1), m.group(2).strip())
            if "ARTÍCULOS Y SECCIONES" not in texto:
                continue
            tipo_centro = texto.splitlines()[2].strip()
            palabras = page.extract_words()
            cols = []
            for i, w in enumerate(palabras):
                if w["text"] == "Sec." and i + 1 < len(palabras):
                    cols.append((palabras[i + 1]["x1"], palabras[i + 1]["text"]))
                elif w["text"] == "TOTAL" and abs(w["top"] - next((p["top"] for p in palabras if p["text"] == "ARTICULOS"), -99)) < 3:
                    cols.append((w["x1"], "TOTAL"))
            # Agrupa palabras en filas con tolerancia (las cifras pueden ir 1 pt más abajo que la etiqueta)
            filas = []
            for w in sorted(palabras, key=lambda w: w["top"]):
                if filas and abs(filas[-1][0]["top"] - w["top"]) < 3:
                    filas[-1].append(w)
                else:
                    filas.append([w])
            for ws in filas:
                ws.sort(key=lambda w: w["x0"])
                art = ws[0]["text"][:3]
                if art not in ("10-", "11-"):
                    continue
                for w in ws[1:]:
                    if re.fullmatch(r"[\d.]+", w["text"]):
                        col = min(cols, key=lambda cx: abs(cx[0] - w["x1"]))[1]
                        if col == "TOTAL":
                            control[(tipo_centro, art[:2])] += num(w["text"])
                        else:
                            c[(col, art[:2], tipo_centro)] += num(w["text"])
    # Comprobación: la suma por secciones debe cuadrar con la columna TOTAL
    for (tc, art), tot in control.items():
        suma = sum(v for (s, a, t), v in c.items() if t == tc and a == art)
        assert abs(suma - tot) < 1, f"Madrid {tc} art{art}: {suma} != {tot}"
    # Las retribuciones de los diputados de la Asamblea van en un artículo propio (17) de la sección 01
    diputados = 0
    with pdfplumber.open(ruta("madrid", "libro02.pdf")) as pdf:
        for page in pdf.pages[:8]:
            for l in (page.extract_text() or "").splitlines():
                m = re.match(r"^17 RETRIBUCIONES DIPUTADOS ASAMBLEA DE MADRID ([\d.]+)$", l.strip())
                if m:
                    diputados += num(m.group(1))
    if not diputados:
        raise ValueError("no se encuentran las retribuciones de los diputados en libro02.pdf")
    c[("01", "10", "ADMINISTRACIÓN GENERAL")] += diputados
    salida = collections.Counter()
    for (sec, art, tipo_centro), v in c.items():
        tipo = None if tipo_centro.startswith("ADMINISTRACI") else "organismo"
        salida[(nombres.get(sec, "Sección " + sec), art, tipo)] += v
    return salida


def desde_json(carpeta, fichero="art10_11.json"):
    """Datos extraídos a mano (o desde el navegador) y guardados en fuentes/<año>/<carpeta>/<fichero>"""
    d = json.loads(ruta(carpeta, fichero).read_text(encoding="utf-8"))
    c = collections.Counter()
    for f in d["filas"]:
        unidad = re.sub(r"^\d+\s+", "", f["seccion"]).strip()
        c[(unidad, f["articulo"], f.get("tipo"))] += f["importe"]
    return c


def baleares():
    """Partidas liquidadas de la CAIB (todos los años en un fichero): crédito 'ordenat' (obligaciones reconocidas)."""
    ws = openpyxl.load_workbook(ruta("baleares", "ppp_dsp_liq.xlsx"), read_only=True).worksheets[0]
    filas = ws.iter_rows(values_only=True)
    cab = list(next(filas))
    i_any, i_sec, i_sector, i_art = cab.index("Exercici"), cab.index("Secció"), cab.index("Sector"), cab.index("Article")
    i_imp = cab.index("Crèdit ordenat")
    c = collections.Counter()
    for r in filas:
        if str(r[i_any]) != str(ANYO) or not r[i_art]:
            continue
        art = str(r[i_art]).split(".")[0]
        if art in ("10", "11"):
            codigo, _, unidad = str(r[i_sec]).partition(".-")
            # Sector CAIB con sección < 50 = Administración General e instituciones; el resto, organismos y entes
            tipo = None if r[i_sector] in ("CAIB", "AGIB") and int(codigo) < 50 else "organismo"
            c[(unidad.strip(), art, tipo)] += float(r[i_imp] or 0)
    if not c:
        raise FaltaFuente(f"el fichero de Baleares aún no trae la liquidación de {ANYO}")
    return c


def conprel_ciudad(nombre):
    """Ceuta y Melilla rinden cuentas como entidades locales: conceptos 100 (órganos de gobierno),
    101 (directivos) y 110 (eventuales) de su liquidación en CONPREL."""
    import pyodbc
    fichero = ruta("conprel", "Liquidaciones.accdb")
    con = pyodbc.connect(r"DRIVER={Microsoft Access Driver (*.mdb, *.accdb)};DBQ=" + str(fichero) + ";ReadOnly=1")
    cur = con.cursor()
    ids = [int(r.idente) for r in cur.execute("SELECT idente FROM tb_inventario WHERE nombreppal LIKE ?", nombre + "%")]
    c = collections.Counter()
    marcas = ",".join("?" * len(ids))
    for r in cur.execute(f"SELECT cdcta, importer FROM tb_economica WHERE tipreig='G' AND cdcta IN ('100   ','101   ','110   ') AND idente IN ({marcas})", *ids):
        cod = r.cdcta.strip()
        if cod == "100":
            c[("Consejo de Gobierno y diputados de la Asamblea (órganos de gobierno)", "10", "gobierno")] += float(r.importer or 0)
        elif cod == "101":
            c[("Personal directivo", "10", "organismo")] += float(r.importer or 0)
        else:
            c[("Personal eventual", "11", "gobierno")] += float(r.importer or 0)
    return c


def pais_vasco():
    nombres = {}
    with open(ruta("paisvasco", "Organikoa_Organica.csv"), encoding="latin-1") as f:
        lector = csv.reader(f, delimiter=";")
        next(lector); next(lector)
        for r in lector:
            if len(r) >= 3 and r[1]:
                nombres[r[1]] = r[2].strip()
    c = collections.Counter()
    with open(ruta("paisvasco", "Datuak_datos.csv"), encoding="latin-1") as f:
        lector = csv.reader(f, delimiter=";")
        next(lector); next(lector)
        for r in lector:
            if len(r) < 15 or r[1] != "1":
                continue
            art = r[10]
            if art in ("10", "11"):
                unidad = nombres.get(r[2] + r[3].zfill(2), "Entidad " + r[2] + " sección " + r[3])
                # Entidad 100 = Administración General; el resto son organismos y entes
                tipo = None if r[2] == "100" else "organismo"
                c[(unidad, art, tipo)] += num(r[14])
    return c


PARLAMENTO = re.compile(r"^(parlamento|parlament|cortes|les corts|corts|asamblea|junta general)", re.I)
ESTATUTARIO = re.compile(
    r"justicia de arag|defensor|procurador del com|s[ií]ndic|sindicatura|c[aá]mara de (cuentas|comptos)|consejo de cuentas"
    r"|audiencia de cuentas|tribunal vasco de cuentas|consejo consultivo|consell consultiu|consejo jur[ií]dico|comisi[oó]n jur[ií]dica"
    r"|comissi[oó] jur[ií]dica|consell de garanties|oficina antifrau|consejo econ[oó]mico y social|consell treball|consejo de transparencia"
    r"|ararteko|valedor|diputado del com[uú]n|acad[eè]mia valenciana|consejo audiovisual|consell de l.audiovisual"
    r"|instituciones propias|consell valenci[aà] de cultura|consello (de contas|da cultura|consultivo)|consell econ[oò]mic|lluita contra la corrupci", re.I)
ORGANISMO = re.compile(
    r"agencia|ag[eè]ncia|^ag\.|instituto|institut|inst\.|\bservicio\b|\bservei\b|consorcio|consorci|hospital|fundaci[oó]n|red de salud|radio|televisi"
    r"|rtva|lanbide|spri|labora|^e\.p\.|entidad|entitat|patronato|centro|centre|escola|escuela|biblioteca|^i\.|aacid|trade|autoridad"
    r"|autoritat|consell catal|academia|osakidetza|^osi |ente p[uú]blico|^c\.estudis", re.I)


def clasificar(unidad):
    if PARLAMENTO.search(unidad):
        return "parlamento"
    if ESTATUTARIO.search(unidad):
        return "estatutario"
    if ORGANISMO.search(unidad):
        return "organismo"
    return "gobierno"


EXTRACTORES = {
    "Andalucía": andalucia,
    "Aragón": aragon,
    "Canarias": canarias,
    "Castilla y León": castilla_leon,
    "Cataluña": cataluna,
    "Navarra": navarra,
    "Comunitat Valenciana": valencia,
    "Región de Murcia": murcia,
    "Comunidad de Madrid": madrid,
    "Galicia": lambda: desde_json("galicia"),
    "Asturias": asturias,
    "La Rioja": lambda: desde_json("larioja"),
    "Cantabria": cantabria,
    "Extremadura": lambda: desde_json("extremadura"),
    "Castilla-La Mancha": castilla_la_mancha,
    "Illes Balears": baleares,
    "Ceuta": lambda: conprel_ciudad("Ceuta"),
    "Melilla": lambda: conprel_ciudad("Melilla"),
    "País Vasco": pais_vasco,
}


def main():
    meta = metadatos()["autonomico"]
    filas, errores = [], {}
    for ccaa, fn in EXTRACTORES.items():
        try:
            datos = fn()
        except (FaltaFuente, ValueError, AssertionError, KeyError, OSError) as e:
            errores[ccaa] = f"{type(e).__name__}: {e}"
            print(f"{ccaa:22s} ERROR {errores[ccaa]}")
            continue
        m = meta[ccaa]
        t = collections.Counter()
        for clave, imp in sorted(datos.items(), key=lambda kv: tuple(str(x) for x in kv[0])):
            if round(imp) == 0:
                continue
            unidad, art = clave[0], clave[1]
            tipo = clave[2] if len(clave) > 2 and clave[2] else clasificar(unidad)
            if tipo == "gobierno" and PARLAMENTO.search(unidad):
                tipo = "parlamento"
            filas.append({"ccaa": ccaa, "unidad": unidad, "tipo": tipo, "articulo": art, "importe": round(imp, 2),
                          "base": m["base"], "fuente": m["texto"], "url": m["url"]})
            t[(tipo, art)] += imp
        print(f"{ccaa:22s} " + "  ".join(f"{k[0][:5]}{k[1]}={v / 1e6:5.1f}" for k, v in sorted(t.items())))
    with open(SALIDA, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["ccaa", "unidad", "tipo", "articulo", "importe", "base", "fuente", "url"])
        w.writeheader()
        w.writerows(filas)
    print("->", SALIDA, len(filas), "filas")
    if errores:
        print(f"\n{len(errores)} comunidades sin datos de {ANYO}: {', '.join(errores)}")
        sys.exit(1)


if __name__ == "__main__":
    main()
